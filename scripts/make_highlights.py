"""Create the journal's short Highlights file using the bundled python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
lines=ROOT.joinpath("highlights.txt").read_text(encoding="utf-8").splitlines()
assert 3 <= len(lines) <= 5 and all(len(s)<=85 for s in lines)
doc=Document()
for border in doc.styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
for name in ("Normal","Title","List Bullet"):
    style=doc.styles[name]
    style.font.name="Times New Roman"
    style.font.underline=False
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(12 if name!="Title" else 16)
doc.add_paragraph("Highlights",style="Title")
for s in lines:
    p=doc.add_paragraph(s,style="List Bullet")
    p.paragraph_format.space_after=Pt(10)
    p.paragraph_format.line_spacing=1.15
doc.core_properties.title="Three-gate evaluation of correction - Highlights"
doc.core_properties.author="Yang Song"
doc.core_properties.last_modified_by="Yang Song"
doc.save(ROOT/"highlights.docx")
print("Highlights characters:", [len(s) for s in lines])
